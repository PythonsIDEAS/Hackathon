from flask import Flask, render_template, request, send_file, flash, redirect, url_for
import os
import tempfile
import requests
import re
import cv2
import easyocr
import numpy as np
import pdfplumber
from docx import Document
from flask_sqlalchemy import SQLAlchemy
from flask_login import UserMixin, login_user, LoginManager, login_required
from flask_bcrypt import Bcrypt


# Initialize EasyOCR reader for English text
reader = easyocr.Reader(['en'])

# Regular expressions to detect common sensitive data patterns
patterns = {
    "IIN": r"\b\d{12}\b",  # INN pattern (12 digits)
    "PHONE_NUMBER": r"\+?\d{1,3}[-\s]?\(?\d{1,4}\)?[-\s]?\d{1,4}[-\s]?\d{1,4}",  # Phone number pattern
    "EMAIL_ADDRESS": r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}",  # Email pattern
    "PERSON": r"\b[A-ZА-ЯЁ][a-zа-яё]+\s[A-ZА-ЯЁ][a-zа-яё]+\b"  # Name pattern
}

# Anonymization placeholders
anonymization_placeholders = {
    "IIN": "[ИИН_АНОНИМНО]",
    "PHONE_NUMBER": "[ТЕЛ_АНОНИМНО]",
    "EMAIL_ADDRESS": "[EMAIL_АНОНИМНО]",
    "PERSON": "[ФИО_АНОНИМНО]"
}

# Function to anonymize text
def anonymize_text(text):
    anonymized_text = text
    try:
        for entity, pattern in patterns.items():
            matches = re.findall(pattern, anonymized_text)
            for match in matches:
                anonymized_text = anonymized_text.replace(match, anonymization_placeholders[entity])
    except Exception as e:
        print(f"An error occurred during anonymization: {e}")
    return anonymized_text

# Function to anonymize image
def anonymize_image(image_path, output_path):
    image = cv2.imread(image_path)
    results = reader.readtext(image_path)
    
    for (bbox, text, prob) in results:
        anon_text = anonymize_text(text)
        if anon_text != text:
            pts = np.array(bbox).astype(int)
            cv2.rectangle(image, tuple(pts[0]), tuple(pts[2]), (0, 0, 0), -1)
            cv2.putText(image, "[АНОНИМНО]", tuple(pts[0]), cv2.FONT_HERSHEY_SIMPLEX, 0.5, (255, 255, 255), 1)
    
    cv2.imwrite(output_path, image)

# Function to anonymize PDF
def anonymize_pdf(input_pdf_path, output_txt_path):
    # Open and read the PDF
    with pdfplumber.open(input_pdf_path) as pdf:
        anonymized_content = ""
        for i, page in enumerate(pdf.pages):
            text = page.extract_text()
            if not text:
                continue

            anonymized_text = anonymize_text(text)
            anonymized_content += f"\n--- Page {i+1} ---\n" + anonymized_text

    # Save anonymized output to a .txt file for now
    with open(output_txt_path, "w", encoding="utf-8") as f:
        f.write(anonymized_content)

# Function to anonymize DOCX
def anonymize_docx(input_docx_path, output_docx_path):
    try:
        doc = Document(input_docx_path)
        for paragraph in doc.paragraphs:
            anonymized_text = anonymize_text(paragraph.text)
            paragraph.text = anonymized_text
        doc.save(output_docx_path)
    except Exception as e:
        print(f"An error occurred: {e}")


app = Flask(__name__)
app.config['SECRET_KEY'] = 'your-secret-key-here'
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///users.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db = SQLAlchemy(app)
bcrypt = Bcrypt(app)
login_manager = LoginManager(app)
login_manager.login_view = 'login'

# Temporary directory to save files
TEMP_DIR = '/tmp/anonymizer_temp/'
RECAPTCHA_SECRET_KEY = '6LeUKD4rAAAAAHzpW94FZS8JAZ59ls9OYa7pU7nC'

def verify_recaptcha(response_token):
    r = requests.post(
        'https://www.google.com/recaptcha/api/siteverify',
        data={'secret': RECAPTCHA_SECRET_KEY, 'response': response_token}
    )
    result = r.json()
    return result.get('success', False)
# Ensure the temporary directory exists
if not os.path.exists(TEMP_DIR):
    os.makedirs(TEMP_DIR)

class User(db.Model, UserMixin):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(150), nullable=False, unique=True)
    email = db.Column(db.String(150), nullable=False, unique=True)
    password = db.Column(db.String(150), nullable=False)

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/about')
def about():
    return render_template('about.html')

@app.route('/features')
def features():
    return render_template('features.html')

@app.route('/contact', methods=['GET', 'POST'])
def contact():
    if request.method == 'POST':
        name = request.form.get('name')
        email = request.form.get('email')
        subject = request.form.get('subject')
        message = request.form.get('message')
        # Here you would typically handle the contact form submission
        # For now, we'll just redirect back with a success message
        flash('Thank you for your message. We will get back to you soon!', 'success')
        return redirect(url_for('contact'))
    return render_template('contact.html')

@app.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'POST':
        username = request.form['username']
        email = request.form['email']
        password = request.form['password']
        confirm_password = request.form['confirm_password']

        # Проверка пароля
        if password != confirm_password:
            flash('Пароли не совпадают', 'danger')
            return redirect(url_for('register'))

        # Проверка, существует ли пользователь
        existing_user = User.query.filter((User.username == username) | (User.email == email)).first()
        if existing_user:
            flash('Пользователь с таким именем или email уже существует', 'warning')
            return redirect(url_for('register'))

        hashed_password = bcrypt.generate_password_hash(password).decode('utf-8')
        new_user = User(username=username, email=email, password=hashed_password)

        db.session.add(new_user)
        db.session.commit()
        flash('Регистрация успешна. Войдите в систему.', 'success')
        return redirect(url_for('login'))

    return render_template('register.html')

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']
        recaptcha_response = request.form.get('g-recaptcha-response')

        # Проверка капчи
        if not verify_recaptcha(recaptcha_response):
            flash('Пожалуйста, подтвердите, что вы не робот', 'danger')
            return redirect(url_for('login'))

        user = User.query.filter_by(username=username).first()
        if user and bcrypt.check_password_hash(user.password, password):
            login_user(user)
            flash('Вы успешно вошли в систему', 'success')
            return redirect(url_for('home'))
        else:
            flash('Неверное имя пользователя или пароль', 'danger')
            return redirect(url_for('login'))

    return render_template('login.html')

@app.route('/text_anonymizer', methods=['GET'])
def text_anonymizer():
    return render_template('text_anonymizer.html')

@app.route('/pdf_anonymizer', methods=['GET'])
def pdf_anonymizer():
    return render_template('pdf_anonymizer.html')

@app.route('/docx_anonymizer', methods=['GET'])
def docx_anonymizer():
    return render_template('docx_anonymizer.html')

@app.route('/anonymize_text', methods=['POST'])
def anonymize_text_route():
    text = request.form['text']
    anonymized_text = anonymize_text(text)
    return render_template('result.html', result=anonymized_text, type='text')

@app.route('/image_anonymizer', methods=['GET'])
def image_anonymizer():
    return render_template('image_anonymizer.html')

@app.route('/anonymize_image', methods=['POST'])
def anonymize_image_route():
    if 'image' not in request.files:
        return 'No image uploaded', 400
    image = request.files['image']
    temp_path = os.path.join(TEMP_DIR, image.filename)
    image.save(temp_path)
    output_path = os.path.join(TEMP_DIR, 'anonymized_' + image.filename)
    anonymize_image(temp_path, output_path)
    return send_file(output_path, as_attachment=True)

@app.route('/anonymize_pdf', methods=['POST'])
def anonymize_pdf_route():
    if 'pdf' not in request.files:
        return 'No PDF uploaded', 400
    pdf = request.files['pdf']
    if not pdf.filename:
        return 'No selected file', 400
    temp_path = os.path.join(TEMP_DIR, pdf.filename)
    pdf.save(temp_path)
    output_path = os.path.join(TEMP_DIR, 'anonymized_' + pdf.filename.replace('.pdf', '.txt'))
    anonymize_pdf(temp_path, output_path)
    return send_file(output_path, as_attachment=True)

@app.route('/anonymize_docx', methods=['POST'])
def anonymize_docx_route():
    if 'docx' not in request.files:
        return 'No DOCX uploaded', 400
    docx = request.files['docx']
    if not docx.filename:
        return 'No selected file', 400
    temp_path = os.path.join(TEMP_DIR, docx.filename)
    docx.save(temp_path)
    output_path = os.path.join(TEMP_DIR, 'anonymized_' + docx.filename)
    anonymize_docx(temp_path, output_path)
    return send_file(output_path, as_attachment=True)

@app.route('/anonymize_database', methods=['POST'])
@login_required
def anonymize_database_route():
    table_name = request.form['table_name']
    columns_to_anonymize = request.form['columns_to_anonymize'].split()
    db_type = request.form['db_type']
    
    try:
        anonymizer = DatabaseAnonymizer(table_name, columns_to_anonymize)
        anonymized_data = anonymizer.anonymize_table()
        return render_template('result.html', result=anonymized_data, type='database')
    except Exception as e:
        return str(e), 400

@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))

if __name__ == '__main__':
    with app.app_context():
        db.create_all()
    app.run(debug=True)